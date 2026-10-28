from pathlib import Path
import re, json, subprocess, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT

sys.stdout.reconfigure(encoding='utf-8')
PROJECT=Path(__file__).resolve().parent
ROOT=PROJECT.parent.parent
OUT=PROJECT
MD=OUT/'REPORT.md'
CONTENT=PROJECT/'report_content.md'
repos={
 'D':('DeepRegular/SmartCut','DeepRegular-SmartCut','84252659110ad34e1af02689f3013badc1561a1e'),
 'P':('skeskinen/smartcut','skeskinen-smartcut','73821ea01c29a4dca81a8865ef7d348d04bb0810'),
 'M':('MediaMolder/mediamolder','mediamolder','9e51f8f761588e2e3c5ef87e9f11dfd5c091ac58'),
 'L':('mifi/lossless-cut','lossless-cut','3ef357d88d51d8a2a36e5bf3c041b3a269bfbc5a'),
 'V':('ozmartian/vidcutter','vidcutter','62d9b0657fd841c74a8d5bb184512ddaddacbce2'),
 'A':('anyc/avcut','avcut','7704c26c95a4540a43e5c9ce9c6470595a78eddc'),
 'X':('fishman/virtualdub','virtualdub-mirror','388a9e98ff08b9b546e4c053c3781a760f5ff976'),
}
spec={
 'D1':('rust/crates/core/src/plan.rs',624),
 'D2':('rust/crates/core/src/leadrefs.rs',1),
 'D3':('rust/crates/core/src/audio.rs',1695),
 'D4':('rust/crates/core/src/cut.rs',667),
 'D5':('docs/technical/rust-core.md',1),
 'D6':('rust/crates/core/src/cut.rs',4727),
 'D7':('rust/crates/vc1/src/lib.rs',1),
 'D8':('rust/crates/core/src/si.rs',1),
 'D9':('rust/crates/core/src/restamp.rs',1),
 'D10':('rust/crates/core/src/verify.rs',1),
 'P1':('smartcut/smart_cut.py',48),
 'P2':('smartcut/video_cutter.py',150),
 'P3':('smartcut/track_cutters.py',35),
 'P4':('smartcut/nal_tools.py',106),
 'P5':('smartcut/media_container.py',80),
 'P6':('README.md',1),
 'M1':('job/handlers_smartcopy.go',1),
 'M2':('job/smartcopy_plan.go',214),
 'M3':('docs/smartcopy.md',1),
 'M4':('job/smartcopy_hevc_ids.go',1),
 'M5':('job/handlers_graph_build.go',379),
 'M6':('job/handlers_smartcopy_cut_test.go',1),
 'M7':('job/handlers_sink.go',58),
 'M8':('job/smartcopy_h264poc.go',1),
 'L1':('src/renderer/src/smartcut.ts',1),
 'L2':('src/renderer/src/hooks/useFfmpegOperations.ts',500),
 'V1':('vidcutter/libs/videoservice.py',261),
 'V2':('vidcutter/libs/config.py',1),
 'A1':('avcut.c',248),
 'A2':('avcut.c',1079),
 'A3':('README.md',1),
 'X3':('VirtualDub/source/DubUtils.cpp',194),
 'X4':('VirtualDub/source/DubProcessVideo.cpp',345),
 'X5':('VirtualDub/source/Dub.cpp',766),
}
refs={}
for key,(path,line) in spec.items():
 repo,folder,sha=repos[key[0]]
 src=ROOT/'sources'/folder/path
 assert src.is_file(),src
 assert line<=len(src.read_text(encoding='utf-8',errors='replace').splitlines()),(path,line)
 refs[key]=f'https://github.com/{repo}/blob/{sha}/{path}#L{line}'
refs['X1']='https://www.virtualdub.org/download.html'
refs['X2']='https://documentation.help/VirtualDub/p-smartrendering.html'
issue_sets={'P':[27,28,23,24],'L':[2494,2050,2049,2013,2912,126],
            'V':[402,258,175,368],'A':[9,19,21,27,25],'D':[9,8],'M':[77,79]}
issue_labels={'P':'smartcut','L':'LosslessCut','V':'VidCutter','A':'avcut','D':'DeepRegular','M':'MediaMolder PR'}
for prefix,numbers in issue_sets.items():
 for n in numbers:
  refs[f'I_{prefix}{n}']=f'https://github.com/{repos[prefix][0]}/{"pull" if prefix=="M" else "issues"}/{n}'

raw=(CONTENT if CONTENT.exists() else MD).read_text(encoding='utf-8').split('\n<!-- REFERENCES -->')[0].rstrip()
raw=raw.replace('它们涉及 Smart Cut 异常、字幕不同步、特定 HEVC 输出损坏和音频问题；',
                '它们分别涉及 H.264 加 AAC 接缝音频拉伸、字幕不同步、特定 HEVC 输出损坏和接缝音频重复；')
raw=raw.replace('对应层次可能是边界计时与重复帧、字幕事件时间、码流参数兼容，以及音频包衔接。',
                '对应层次可能是音频包时间与 concat 偏移、字幕事件时间、码流参数兼容，以及音频包衔接。')
raw=raw.replace('主要独有设计','突出设计')
raw=re.sub(r'(\[[A-Z][A-Z_0-9]*\])\1+',r'\1',raw)
raw=raw.replace('还涉及图与 sink 的时间逻辑。[M5][I_M79]','还涉及图与 sink 的时间逻辑。[M5][M7][I_M79]')
raw=raw.replace('- [M6] handlers_smartcopy_cut_test.go 的回归夹具。','- [M6] handlers_smartcopy_cut_test.go 的回归夹具。\n- [M7] handlers_sink.go 的已裁剪音频时间处理。\n- [M8] smartcopy_h264poc.go 的 H.264 接缝头部重写。') if '- [M7]' not in raw else raw
found=set(re.findall(r'\[([A-Z][A-Z_0-9]*)\]',raw))
assert found <= refs.keys(), found-refs.keys()
CONTENT.write_text(raw+'\n',encoding='utf-8')
def md_reference(match):
 key=match[1]
 if key.startswith('I_'):
  m=re.fullmatch(r'I_([A-Z])(\d+)',key)
  label=f'{issue_labels[m[1]]} #{m[2]}'
 else:label=key
 return f' [{label}]({refs[key]})'
MD.write_text(re.sub(r'\[([A-Z][A-Z_0-9]*)\]',md_reference,raw)+'\n',encoding='utf-8')

manifest=[]
for prefix,(repo,folder,sha) in repos.items():
 repo_path=str(ROOT/'sources'/folder)
 actual=subprocess.check_output(['git','-c','safe.directory='+repo_path.replace('\\','/'),'-C',repo_path,'rev-parse','HEAD'],text=True).strip()
 assert actual==sha,(repo,actual,sha)
 manifest.append({'repository':repo,'sha':sha,'local_source':str(ROOT/'sources'/folder),'as_of':'2026-10-25'})
(PROJECT/'evidence'/'source_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
(PROJECT/'evidence'/'report_references.json').write_text(json.dumps(refs,ensure_ascii=False,indent=2),encoding='utf-8')

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.left_margin=sec.right_margin=Inches(.8)
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.7)
sec.header_distance=sec.footer_distance=Inches(.3)
sec.different_first_page_header_footer=False
for e in list(doc.settings.element.findall(qn('w:evenAndOddHeaders'))):doc.settings.element.remove(e)
for s in doc.styles:
 for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for fonts in s.element.iter(qn('w:rFonts')):
  for attr in list(fonts.attrib):
   if 'Theme' in attr or 'theme' in attr:del fonts.attrib[attr]
style_sizes={'Normal':11,'Title':23,'Subtitle':12,'Heading 1':16,'Heading 2':12.5,'Heading 3':11.5,'List Bullet':10.5}
for name,size in style_sizes.items():
 s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
 s.font.italic=False
 s.paragraph_format.line_spacing=1.22
 s.paragraph_format.space_after=Pt(6)
 if name.startswith('Heading'):
  s.font.bold=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)
  s.paragraph_format.keep_with_next=True
 if name=='Title':
  s.paragraph_format.space_after=Pt(8)
 if name=='List Bullet':
  s.paragraph_format.space_after=Pt(3)
  s.paragraph_format.left_indent=Inches(.16)
  s.paragraph_format.first_line_indent=Inches(-.13)
doc.styles['Normal'].paragraph_format.widow_control=True
doc.core_properties.title='视频 Smart Rendering 源码研究报告'
doc.core_properties.subject='七个项目的视频 音频 容器实现与 issue 对照'
doc.core_properties.author='Codex'
doc.core_properties.keywords='Smart Rendering, GOP, H.264, HEVC, audio, container'

hp=sec.header.paragraphs[0]
hr=hp.add_run('Smart Rendering 源码研究');hr.font.size=Pt(8);hr.font.color.rgb=RGBColor(0,0,0)
hr._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
hp.paragraph_format.space_after=Pt(0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fr=fp.add_run('2026 年 10 月 26 日    ·    ');fr.font.size=Pt(8)
fr._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)

def add_hyperlink(p,label,url,small=True):
 h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
 r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Calibri');fonts.set(qn('w:eastAsia'),'Microsoft YaHei');pr.append(fonts)
 color=OxmlElement('w:color');color.set(qn('w:val'),'225B86');pr.append(color)
 if small:
  sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');pr.append(sz)
 r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)

def inline(p,text):
 # Input contains only plain text, bold spans, and verified reference tokens.
 bold=False
 for token in re.split(r'(\*\*|\[[A-Z][A-Z_0-9]*\])',text):
  if token=='**':bold=not bold;continue
  if re.fullmatch(r'\[[A-Z][A-Z_0-9]*\]',token):
   key=token[1:-1]
   if key.startswith('I_'):
    match=re.fullmatch(r'I_([A-Z])(\d+)',key)
    label=f'〔{issue_labels[match[1]]} #{match[2]}〕'
   else:label=f'〔{key}〕'
   add_hyperlink(p,label,refs[key]);continue
  r=p.add_run(token);r.bold=True if bold else None

def table(lines):
 rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
 rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',x) for x in r)]
 n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 if n==4:widths=[1.10,1.60,2.35,1.85]
 elif rows[0][1]=='本次源码快照':widths=[1.38,2.25,3.27]
 else:widths=[1.15,2.75,3.0]
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,rowdata in enumerate(rows):
  row=t.rows[0] if i==0 else t.add_row()
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:
   repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
  for j,(cell,content) in enumerate(zip(row.cells,rowdata)):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr()
   mar=OxmlElement('w:tcMar')
   for side,val in [('top','90'),('bottom','90'),('left','105'),('right','105')]:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   if i==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EFF5');pr.append(shade)
   p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.15
   inline(p,content)
   for r in p.runs:r.font.size=Pt(10);r.bold=(i==0)
 borders=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
 t._tbl.tblPr.append(borders)
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
 p.add_run().font.size=Pt(2)

lines=raw.splitlines();i=0;nonempty=0
while i<len(lines):
 line=lines[i]
 if not line.strip():i+=1;continue
 if line.startswith('|'):
  batch=[]
  while i<len(lines) and lines[i].startswith('|'):batch.append(lines[i]);i+=1
  table(batch);continue
 if line.startswith('# '):
  p=doc.add_paragraph(style='Title');inline(p,line[2:])
 elif line.startswith('## '):
  p=doc.add_paragraph(style='Heading 1');inline(p,line[3:])
 elif line.startswith('### '):
  p=doc.add_paragraph(style='Heading 2');inline(p,line[4:])
 elif line.startswith('- '):
  p=doc.add_paragraph(style='List Bullet');inline(p,line[2:])
 elif nonempty==1:
  p=doc.add_paragraph(style='Subtitle');inline(p,line)
 elif line.startswith('交付日期：'):
  p=doc.add_paragraph();inline(p,line)
  for r in p.runs:r.font.size=Pt(9)
 else:
  p=doc.add_paragraph();inline(p,line)
 i+=1;nonempty+=1

path=OUT/'REPORT.docx';doc.save(path)
print(json.dumps({'docx':str(path),'markdown':str(MD),'characters':len(raw),'source_links':len(spec),'total_references':len(refs),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables)},ensure_ascii=False))
